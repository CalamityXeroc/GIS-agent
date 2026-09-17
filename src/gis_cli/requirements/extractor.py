# -*- coding: utf-8 -*-
"""Extract structured requirements from a task goal and/or documents.

This turns "普通人给的一份任务书" into a checkable requirement list the
verifier can grade against — the missing piece between natural language and
acceptance criteria.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from ..engine.tool_codec import extract_json_object
from . import Requirement

logger = logging.getLogger(__name__)


def read_document_text(path: str | Path, max_chars: int = 24000) -> str:
    """Extract plain text from docx/pdf/txt (shared by tool + extractor)."""
    target = Path(path)
    suffix = target.suffix.lower()
    if suffix == ".docx":
        import docx  # type: ignore

        doc = docx.Document(str(target))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    elif suffix == ".pdf":
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(target))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    else:
        text = target.read_text(encoding="utf-8", errors="replace")
    return text[:max_chars]


class RequirementExtractor:
    """LLM-backed requirement extraction with document support."""

    def __init__(self, llm_client: Any, max_doc_chars: int = 24000):
        self.llm_client = llm_client
        self.max_doc_chars = max_doc_chars

    def extract(
        self,
        goal: str,
        doc_paths: list[str] | None = None,
        catalog_digest: str = "",
    ) -> list[Requirement]:
        """Extract requirements; returns [] when extraction is not possible."""
        docs_block = ""
        for path in doc_paths or []:
            try:
                text = read_document_text(path, self.max_doc_chars)
            except Exception as exc:
                logger.warning("doc read failed %s: %s", path, exc)
                continue
            docs_block += f"\n\n### 文档：{Path(path).name}\n{text[:self.max_doc_chars]}"

        prompt = (
            "你是 GIS 需求分析专家。请把下面的任务目标（和任务书节选，如有）提炼成一份可验收的需求清单。\n"
            "每条需求必须：具体、可检验、带明确的验收标准（数据/字段/坐标系/数量/格式等可核对对象）。\n"
            "不要把“说明性/背景性”内容写成需求；不要编造数据里不存在的字段名。\n"
            "只输出一个 JSON 对象：\n"
            '{"requirements": [{"id": "r1", "text": "...", "acceptance": "..."}]}\n\n'
            f"## 任务目标\n{goal}\n"
        )
        if docs_block:
            prompt += f"## 任务书节选{docs_block[:self.max_doc_chars]}\n"
        if catalog_digest:
            prompt += f"\n## 当前数据目录（供参考，字段以此为准）\n{catalog_digest[:4000]}\n"

        try:
            response = self.llm_client.chat(
                [
                    {"role": "system", "content": "你只输出一个 JSON 对象。"},
                    {"role": "user", "content": prompt},
                ],
                task_type="requirements",
                temperature=0.1,
            )
        except Exception as exc:
            logger.warning("requirement extraction failed: %s", exc)
            return []

        payload = extract_json_object(response.content)
        if not isinstance(payload, dict):
            return []
        items = payload.get("requirements") or []
        out: list[Requirement] = []
        for i, raw in enumerate(items, start=1):
            if not isinstance(raw, dict):
                continue
            text = str(raw.get("text", "") or "").strip()
            if not text:
                continue
            out.append(
                Requirement(
                    id=str(raw.get("id", "") or f"r{i}"),
                    text=text,
                    acceptance=str(raw.get("acceptance", "") or ""),
                    source="extracted",
                )
            )
        return out
